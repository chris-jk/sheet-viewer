"""Write Sheet Viewer's export as an Excel workbook. Reads JSON on stdin:

{"path": ..., "sheets": [{"name": ..., "header": bool, "rows": [[text, ...], ...], "links": [{"r", "c", "u"}, ...],
                          "styles": [{"f", "c", "b", "i"}, ...], "styled": [row, column, style, ...]}]}

Text that reads as a number, a dollar amount, a percent or a date goes in as
one, so sums and sorts work in Excel. Everything else goes in as text, even
when it starts with "=": formulas are never written, because the viewer only
ever had their values.
"""
import json
import os
import re
import sys
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

NUMBER = re.compile(r"(-?)(\$?)(\d{1,3}(?:,\d{3})+|\d+)?(\.\d+)?(%?)")
LINK_BLUE = "0563C1"


def typed(text):
    """A cell's text as (value, number format): figures and dates become real ones."""
    match = NUMBER.fullmatch(text)
    if match and (match[3] or match[4]):
        sign, dollar, whole, fraction, percent = match.groups()
        digits = (whole or "0").replace(",", "")
        significant = len((digits + (fraction or ".")[1:]).lstrip("0"))
        bare = not (dollar or percent or fraction or "," in (whole or ""))
        # A leading zero is a code (a ZIP, a part number), not a quantity. So is a long bare run of digits:
        # Excel keeps 15 digits of any number and would show a card or tracking number as 4.11E+15.
        if not (len(digits) > 1 and digits[0] == "0") and significant <= 15 and not (bare and len(digits) > 11):
            value = float(digits + fraction) if fraction else int(digits)
            if percent:
                value = value / 100
            if sign:
                value = -value
            grouped = "," in (whole or "") or bool(dollar)
            form = ('"$"' if dollar else "") + ("#,##0" if grouped else "0")
            form += ("." + "0" * (len(fraction) - 1) if fraction else "") + percent
            return value, None if form == "0" else form
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return date.fromisoformat(text), "yyyy-mm-dd"
        if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", text):
            return datetime.fromisoformat(text), "yyyy-mm-dd hh:mm:ss"
    except ValueError:
        pass
    return text, None


def page_names(names):
    """Tab names Excel accepts: 31 characters, none of []:*?/\\, and no two alike whatever their case."""
    taken, out = set(), []
    for name in names:
        base = re.sub(r"[\[\]:*?/\\]", "-", name).strip().strip("'")[:31].rstrip() or "Sheet"
        candidate, count = base, 1
        while candidate.lower() in taken:
            count += 1
            candidate = base[: 31 - len(f" {count}")] + f" {count}"
        taken.add(candidate.lower())
        out.append(candidate)
    return out


def write(book):
    out = Workbook()
    out.remove(out.active)
    for sheet, name in zip(book["sheets"], page_names([sheet["name"] for sheet in book["sheets"]])):
        page = out.create_sheet(name)
        widths = {}
        for r, row in enumerate(sheet["rows"], start=1):
            for c, text in enumerate(row, start=1):
                text = ILLEGAL_CHARACTERS_RE.sub("", text)  # one stray control character would fail the whole file
                if text == "":
                    continue
                value, form = typed(text)
                cell = page.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # or text starting with "=" would be stored as a live formula
                if form:
                    cell.number_format = form
                if r <= 500:
                    widths[c] = max(widths.get(c, 0), len(text))
        fonts = {}
        styled = sheet.get("styled", [])
        for r, c, s in zip(styled[0::3], styled[1::3], styled[2::3]):
            look = sheet["styles"][s]
            cell = page.cell(row=r + 1, column=c + 1)
            if look.get("f"):
                cell.fill = PatternFill("solid", fgColor=look["f"])
            fonts[(r, c)] = {"bold": bool(look.get("b")), "italic": bool(look.get("i")), "color": look.get("c")}
        for link in sheet.get("links", []):
            page.cell(row=link["r"] + 1, column=link["c"] + 1).hyperlink = link["u"]
            look = fonts.setdefault((link["r"], link["c"]), {})
            look["underline"] = "single"
            look["color"] = look.get("color") or LINK_BLUE
        if sheet.get("header"):
            for c in range(len(sheet["rows"][0]) if sheet["rows"] else 0):
                fonts.setdefault((0, c), {}).setdefault("bold", True)
            page.freeze_panes = "A2"
        for (r, c), look in fonts.items():
            page.cell(row=r + 1, column=c + 1).font = Font(**{key: value for key, value in look.items() if value})
        for c, width in widths.items():
            page.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 6), 60)
    if not out.worksheets:
        out.create_sheet("Sheet")
    partial = book["path"] + ".part"
    out.save(partial)
    os.replace(partial, book["path"])


if __name__ == "__main__":
    try:
        write(json.load(sys.stdin))
    except Exception as error:
        raise SystemExit(f"Couldn't write the workbook: {error or type(error).__name__}")
