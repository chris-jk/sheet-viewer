"""Read an Excel or Numbers file for Sheet Viewer and print it as JSON.

{"sheets": [{"name": ..., "rows": [[text, ...], ...], "links": [{"r": row, "c": column, "u": address}, ...],
             "styles": [{"f": fill, "c": text color, "b": bold, "i": italic}, ...], "styled": [row, column, style, ...]}]}

Rows keep their place (row 0 is the file's row 1) and every cell is the text
the spreadsheet would show: a formula comes through as its last saved value,
or as the formula itself when the file holds no value for it.
"styled" is a flat run of triples naming the cells that are not plain.
CSV never gets here; the app reads that itself.
"""
import colorsys
import json
import os
import re
import sys
import zipfile
from datetime import datetime, time
from decimal import ROUND_HALF_UP, Decimal
from xml.etree import ElementTree

PLAIN_NUMBER = re.compile(r"^(\$?)[#0,]*0(\.0+)?(%?)$")
# A formula cell that ends without a value after it.
UNVALUED_FORMULA = re.compile(rb"(?:</f>|<f[^>]*/>)\s*(?:<v\s*/>|<v></v>)?\s*</c>")


def value_text(value):
    """Text for a bare value that carries no display format."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time.min else value.isoformat(sep=" ")
    if isinstance(value, float):
        return f"{value:.15g}"
    return str(value)


def excel_text(cell):
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Honour plain money, percent and fixed-decimal formats; anything else prints raw.
        fmt = re.sub(r"\[\$([^\]-]*)[^\]]*\]", r"\1", (cell.number_format or "").split(";")[0])
        fmt = re.sub(r'[_*].|\[[^\]]*\]|["\\\s()]', "", fmt)
        match = PLAIN_NUMBER.match(fmt)
        if match:
            dollar, decimals, percent = match.groups()
            places = len(decimals or ".") - 1
            try:
                # Excel rounds a half away from zero, on the number as it reads in decimal. Python's float
                # formatting rounds the binary value half to even: 2.5 would show as 2 and 0.125 as 0.12.
                exact = Decimal(repr(value)) * (100 if percent else 1)
                text = f"{exact.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP):{',' if ',' in fmt else ''}f}"
            except ArithmeticError:
                return value_text(value)
            sign = "-" if text.startswith("-") else ""
            return sign + dollar + text.lstrip("-") + percent
    return value_text(value)


def numbers_text(cell):
    if cell.value is None:
        return ""
    try:
        text = cell.formatted_value
    except Exception:
        text = None
    return value_text(cell.value) if text in (None, "", str(cell.value)) else text


def theme_colors(book):
    """The workbook theme's colors, in the order cell styles number them."""
    drawing = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    try:
        scheme = ElementTree.fromstring(book.loaded_theme).find(f".//{drawing}clrScheme")
    except Exception:
        return []
    found = {}
    for entry in scheme if scheme is not None else []:
        if len(entry):
            found[entry.tag.replace(drawing, "")] = entry[0].get("val" if entry[0].tag.endswith("srgbClr") else "lastClr")
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
    return [found.get(name) or "000000" for name in order]


def color_hex(color, theme):
    """A cell color as RRGGBB, or None when it is left to the app (automatic)."""
    from openpyxl.styles.colors import COLOR_INDEX

    try:
        if color is None:
            return None
        if color.type == "rgb":
            value = color.rgb[-6:]
        elif color.type == "theme":
            value = theme[color.theme]
        elif color.type == "indexed" and color.indexed < len(COLOR_INDEX):
            value = COLOR_INDEX[color.indexed][-6:]
        else:
            return None
        red, green, blue = (int(value[i : i + 2], 16) / 255 for i in (0, 2, 4))
        tint = color.tint or 0
        if tint:  # Excel lightens or darkens a theme color by a fraction.
            hue, light, saturation = colorsys.rgb_to_hls(red, green, blue)
            light = light * (1 + tint) if tint < 0 else light + (1 - light) * tint
            red, green, blue = colorsys.hls_to_rgb(hue, light, saturation)
        return "".join(f"{round(part * 255):02X}" for part in (red, green, blue))
    except Exception:
        return None


def has_unvalued_formulas(path):
    """True when a formula cell holds no saved result: a workbook a script wrote and no spreadsheet app has opened."""
    try:
        with zipfile.ZipFile(path) as archive:
            pages = [name for name in archive.namelist() if name.startswith("xl/worksheets/") and name.endswith(".xml")]
            return any(UNVALUED_FORMULA.search(archive.read(name)) for name in pages)
    except Exception:
        return False


def read_xlsx(path):
    from openpyxl import load_workbook

    # A full load, not read_only: only that keeps the links and styles on cells.
    book = load_workbook(path, data_only=True)
    # Values and formulas load separately. The second load is paid only by a file that needs it.
    written = load_workbook(path).worksheets if has_unvalued_formulas(path) else None
    theme = theme_colors(book)
    sheets = []
    for position, sheet in enumerate(book.worksheets):
        rows, links, styles, styled, seen = [], [], [], [], {}
        formulas = written[position].iter_rows() if written else None
        for r, row in enumerate(sheet.iter_rows()):
            rows.append([excel_text(cell) for cell in row])
            if formulas is not None:
                # No saved value: show the formula, so a missing total reads as missing and not as empty.
                for c, source in enumerate(next(formulas, ())):
                    if c < len(row) and rows[r][c] == "" and source.data_type == "f":
                        rows[r][c] = str(getattr(source.value, "text", source.value))
            for c, cell in enumerate(row):
                if cell.hyperlink is not None and cell.hyperlink.target:
                    links.append({"r": r, "c": c, "u": cell.hyperlink.target})
                if not cell.has_style:
                    continue
                # Thousands of cells share a handful of styles: work each one out once.
                if cell.style_id not in seen:
                    look = {
                        "f": color_hex(cell.fill.fgColor, theme) if cell.fill.fill_type == "solid" else None,
                        "c": color_hex(cell.font.color, theme),
                        "b": bool(cell.font.b),
                        "i": bool(cell.font.i),
                    }
                    if look["c"] == "000000":  # Black is what unstyled text is anyway.
                        look["c"] = None
                    if not any(look.values()):
                        seen[cell.style_id] = None
                    elif look in styles:
                        seen[cell.style_id] = styles.index(look)
                    else:
                        seen[cell.style_id] = len(styles)
                        styles.append(look)
                if seen[cell.style_id] is not None:
                    styled += [r, c, seen[cell.style_id]]
        sheets.append({"name": sheet.title, "rows": rows, "links": links, "styles": styles, "styled": styled})
    return sheets


def read_calamine(path):
    """The older and binary formats (.xls, .xlsb, .ods): values only, no display formats or links."""
    from python_calamine import CalamineWorkbook

    book = CalamineWorkbook.from_path(path)
    return [
        {
            "name": name,
            "rows": [[value_text(v) for v in row] for row in book.get_sheet_by_name(name).to_python(skip_empty_area=False)],
            "links": [],
            "styles": [],
            "styled": [],
        }
        for name in book.sheet_names
    ]


def read_numbers(path):
    from numbers_parser import Document

    sheets = []
    for sheet in Document(path).sheets:
        for table in sheet.tables:
            name = sheet.name if len(sheet.tables) == 1 else f"{sheet.name} - {table.name}"
            rows = [[numbers_text(cell) for cell in row] for row in table.rows()]
            sheets.append({"name": name, "rows": rows, "links": [], "styles": [], "styled": []})
    return sheets


READERS = {
    ".xlsx": read_xlsx,
    ".xlsm": read_xlsx,
    ".xltx": read_xlsx,
    ".xltm": read_xlsx,
    ".xls": read_calamine,
    ".xlsb": read_calamine,
    ".ods": read_calamine,
    ".numbers": read_numbers,
}


def trimmed(rows):
    """Drop the blank rows and columns sheets end with."""
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
    return [r[:width] for r in rows]


def read(path):
    path = path.rstrip("/")
    name = os.path.basename(path)
    reader = READERS.get(os.path.splitext(path)[1].lower())
    if reader is None:
        raise SystemExit(f"Sheet Viewer can't read {name}")
    try:
        sheets = reader(path)
    except Exception as error:  # Locked, damaged, or not the format its name claims.
        raise SystemExit(f"Couldn't read {name}: {error or type(error).__name__}")
    sheets = [dict(sheet, rows=trimmed(sheet["rows"])) for sheet in sheets]
    # Excel pads workbooks with blank sheets.
    return {"sheets": [sheet for sheet in sheets if sheet["rows"]]}


if __name__ == "__main__":
    book = read(sys.argv[1])
    sys.stdout.buffer.write(json.dumps(book, ensure_ascii=False, separators=(",", ":")).encode("utf-8", "replace"))
